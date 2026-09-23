import io
import json
from docx import Document
from pypdf import PdfReader
from app.topic_analysis import topic_blocks, extraction_windows, analyze_topics
from app.schemas import Protocol
from app.export import docx_bytes, pdf_bytes
import pytest


def sample():
    return dict(title='Две темы',meeting_date='2026-09-23',speakers={},participants=[],segments=[
        dict(id=0,speaker='A',text='Первая тема — логистика. Загрузка склада 70 процентов.'),
        dict(id=1,speaker='A',text='Первое: Дана, подготовь план до пятницы. Второе: Болат, проверь договор к среде.'),
        dict(id=2,speaker='A',text='Подготовку плана обсудили, сроки и исполнителей согласовали. По логистике всё, переходим ко второму вопросу, безопасность.'),
        dict(id=3,speaker='A',text='Проверка датчиков не выполнена. Требуется аудит.'),
        dict(id=4,speaker='A',text='Дана, проверь датчики.'),
        dict(id=5,speaker='A',text='Срок проверки — до понедельника.')])


def test_explicit_transition_overrides_overfragmented_outline():
    m=sample()
    blocks=topic_blocks(json.dumps({'topics':[dict(title='Логистика',start_source_id=0),dict(title='План',start_source_id=1),dict(title='Договор',start_source_id=2),dict(title='Безопасность',start_source_id=3)]}),m)
    assert [title for title,_ in blocks]==['Логистика','Безопасность']
    assert [r['id'] for r in blocks[0][1]]==[0,1,2]
    assert [r['id'] for r in blocks[1][1]]==[3,4,5]


def test_windows_cover_every_row_once_and_keep_neighbor_context():
    rows=[dict(id=i,text='Коллеги обсудили бюджет проекта, риски поставки сырья и план работы отдела на следующий месяц.') for i in range(9)]
    windows=list(extraction_windows(rows,limit=250))
    assert [r['id'] for core,_ in windows for r in core]==list(range(9))
    assert windows[1][1][0]['id']==0
    assert windows[1][1][-1]['id']==6


def test_grouped_analysis_and_exports_keep_all_tasks():
    m=sample()
    calls=iter([
        {'topics':[dict(title='Логистика',start_source_id=0),dict(title='Безопасность',start_source_id=3)]},
        {'summary_facts':[dict(text='Загрузка склада — 70 процентов.',source_ids=[0],evidence='Загрузка склада 70 процентов.')], 'tasks':[
            dict(title='Подготовить план',owner='Дана',deadline_text='до пятницы',source_ids=[1],evidence='Дана, подготовь план до пятницы.'),
            dict(title='Проверить договор',owner='Болат',deadline_text='к среде',source_ids=[1],evidence='Болат, проверь договор к среде.')]},
        {'summary_facts':[dict(text='Проверка датчиков не выполнена.',source_ids=[3],evidence='Проверка датчиков не выполнена.')], 'tasks':[
            dict(title='Проверить датчики',owner='Дана',deadline_text='до понедельника',source_ids=[4],evidence='Дана, проверь датчики.',context_evidence=[dict(source_id=5,text='Срок проверки — до понедельника.')])]}])
    p,_=analyze_topics(m,lambda *args: json.dumps(next(calls),ensure_ascii=False))
    assert len(p['tasks'])==3
    assert [t['topic_id'] for t in p['tasks']]==['topic-1','topic-1','topic-2']
    assert p['tasks'][2]['deadline_text']=='до понедельника'
    m['protocol']=p
    doc=Document(io.BytesIO(docx_bytes(m)))
    assert len(doc.tables)==2
    assert len(doc.tables[0].rows)==3 and len(doc.tables[1].rows)==2
    assert 'Тема 1 — Логистика' in [p.text for p in doc.paragraphs]
    text='\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(pdf_bytes(m))).pages)
    assert 'Тема 2 — Безопасность' in text
    assert 'Проверить договор' in text and 'Проверить датчики' in text
    with pytest.raises(ValueError): Protocol(topics=p['topics'],tasks=[{**p['tasks'][0],'topic_id':'missing'}])


def test_bad_generation_fails_instead_of_silently_dropping_tasks():
    responses=iter(['{"topics":[]}', '{"tasks": ['])
    with pytest.raises(ValueError): analyze_topics(sample(),lambda *args:next(responses))


def test_grounding_recovers_source_case_row_and_explicit_deadline():
    from app.analysis import parse_result, ground_protocol
    m=sample()
    m['segments'][1]['text']='Первое — разработать единый план поставки сырья для всех подразделений группы, ответственной Дана Сериковна, срок до пятнадцатого октября.'
    draft=parse_result(json.dumps({'tasks':[dict(title='Разработать единый план поставки сырья',owner='Дана Сериковна',deadline_text='15 октября',source_ids=[0],evidence='первое — разработать единый план поставки сырья для всех подразделений группы, ответственной Дана Сериковна, срок до пятнадцатого октября')]}))
    p,_=ground_protocol(draft,m)
    assert len(p['tasks'])==1
    task=p['tasks'][0]
    assert task['source_ids']==[1]
    assert task['evidence'] in m['segments'][1]['text']
    assert task['deadline_text']=='до пятнадцатого октября'
    assert task['owner']=='Дана Сериковна'


def test_long_action_prefix_restores_actual_source_not_model_name():
    from app.analysis import parse_result, ground_protocol
    m=sample()
    m['segments'][1]['text']='Второе — провести совещание с проектной командой и зафиксировать график поставки оборудования, ответственный Айнұр Сериковна, срок до двадцатого сентября.'
    draft=parse_result(json.dumps({'tasks':[dict(title='Провести совещание с проектной командой',owner='Айнур Сериковна',source_ids=[1],evidence='второе — провести совещание с проектной командой и зафиксировать график поставок, ответственный Айнур Сериковна, срок до двадцатого сентября')]}))
    p,_=ground_protocol(draft,m)
    task=p['tasks'][0]
    assert task['owner']=='Айнұр Сериковна'
    assert task['evidence'] in m['segments'][1]['text']
    assert task['deadline_text']=='до двадцатого сентября'


def test_adjacent_confirmation_merges_deadline_but_not_extra_task():
    from app.topic_analysis import consolidate_tasks
    rows=[dict(id=0,text='Дана, запросите заключение юристов.'),dict(id=1,text='Хорошо, запрошу заключение, к среде будет ответ.')]
    tasks=[dict(title='Запросить заключение юристов',topic_id='t',owner='Дана',deadline_text=None,source_ids=[0],evidence=rows[0]['text'],context_evidence=[]),
           dict(title='Получить заключение',topic_id='t',owner=None,deadline_text='к среде',source_ids=[1],evidence=rows[1]['text'],context_evidence=[])]
    result=consolidate_tasks(tasks,rows)
    assert len(result)==1 and result[0]['deadline_text']=='к среде'
    assert result[0]['owner']=='Дана' and result[0]['source_ids']==[0,1]
    assert result[0]['context_evidence'][0]['text']==rows[1]['text']


def test_named_address_before_instruction_is_cited():
    from app.analysis import ground_protocol, parse_result
    m=sample()
    m['segments']=[dict(id=0,speaker='A',text='Дана Сериковна, вы отвечаете за проверку?',uncertain=False),
                   dict(id=1,speaker='B',text='Да.',uncertain=False),
                   dict(id=2,speaker='A',text='Проверьте оборудование до пятницы.',uncertain=False)]
    draft=parse_result(json.dumps({'tasks':[dict(title='Проверить оборудование',owner='Дана Сериковна',source_ids=[2],evidence=m['segments'][2]['text'])]}))
    p,_=ground_protocol(draft,m)
    assert p['tasks'][0]['owner']=='Дана Сериковна'
    assert p['tasks'][0]['deadline_text']=='до пятницы'
    assert any(q['source_id']==0 for q in p['tasks'][0]['context_evidence'])
