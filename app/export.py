from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape


def clock(seconds):
    seconds = int(seconds)
    return f"{seconds // 60:02d}:{seconds % 60:02d}"


def contents(m):
    p = m.get("protocol") or {}
    state = (
        "Подтверждён пользователем"
        if p.get("approved")
        else "Черновик — требуется проверка"
    )
    yield "title", m["title"]
    yield "text", f"Дата: {m['meeting_date']} · {state}"
    yield "text", "Участники: " + (", ".join(m.get("participants", [])) or "Не указаны")
    yield "heading", "Краткое содержание"
    yield "text", p.get("summary") or "Саммари не подготовлено."
    yield "heading", "Решения"
    for decision in p.get("decisions", []):
        yield "text", "• " + decision
    if not p.get("decisions"):
        yield "text", "Решения не зафиксированы."
    yield "heading", "Поручения"
    for i, t in enumerate(p.get("tasks", []), 1):
        yield "text", f"{i}. {t['title']}"
        yield (
            "text",
            f"Ответственный: {t.get('owner') or 'Не указан'}; срок: {t.get('due_date') or t.get('deadline_text') or 'Не указан'}; статус: {dict(open='К выполнению', in_progress='В работе', done='Выполнено').get(t.get('status'), 'К выполнению')}",
        )
        if t.get("evidence"):
            yield "text", "Основание: «" + t["evidence"] + "»"
    if not p.get("tasks"):
        yield "text", "Поручения не зафиксированы."
    yield "heading", "Транскрипт"
    for s in m.get("segments", []):
        name = m.get("speakers", {}).get(s["speaker"], s["speaker"])
        yield "text", f"[{clock(s['start'])}] {name}: {s['text']}"


def docx_bytes(m):
    from docx import Document
    from docx.shared import Pt, RGBColor, Mm
    from docx.oxml.ns import qn

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(18)
    for element in doc.styles.element.iter():
        for child in list(element):
            if child.tag == qn("w:pBdr"):
                element.remove(child)
    for style in ("Normal", "Title", "Heading 1"):
        doc.styles[style].font.name = "Noto Sans"
        doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(13)
    for kind, text in contents(m):
        if kind == "title":
            doc.add_paragraph(text, "Title")
        elif kind == "heading":
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def pdf_bytes(m):
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.pagesizes import A4

    font = "HackAlemNoto"
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(
            TTFont(font, str(Path(__file__).parent / "fonts/NotoSans-Regular.ttf"))
        )
    styles = {
        kind: ParagraphStyle(
            kind,
            fontName=font,
            fontSize=size,
            leading=size * 1.5,
            spaceAfter=8,
            keepWithNext=kind in ("title", "heading"),
        )
        for kind, size in [("title", 20), ("heading", 13), ("text", 10)]
    }
    buffer = BytesIO()
    story = [
        Paragraph(escape(text).replace("\n", "<br/>"), styles[kind])
        for kind, text in contents(m)
    ]

    def footer(canvas, doc):
        canvas.setFont(font, 8)
        canvas.drawRightString(A4[0] - 42, 24, str(doc.page))

    SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=42,
        leftMargin=42,
        topMargin=40,
        bottomMargin=42,
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
